import os
from typing import List, Dict, Any, Optional
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tools.citation_formatter import build_academic_reference_data


from tools.paths import STORAGE_DIR


def set_a4_margins(doc: Document):
    """Menyetel ukuran kertas A4 dan batas tepi standar akademis."""
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(4.0)
        section.left_margin = Cm(4.0)
        section.bottom_margin = Cm(3.0)
        section.right_margin = Cm(3.0)


def create_assignment_docx(
    title: str,
    sections: List[Dict[str, Any]],
    references: List[Dict[str, Any]],
    output_filename: str = "Tugas_Kuliah.docx",
    language: Optional[str] = None,
    identity_lines: Optional[List[str]] = None
) -> str:
    """Merakit naskah tugas lengkap menjadi berkas Word docx format A4 rapi."""
    os.makedirs(STORAGE_DIR, exist_ok=True)
    file_path = os.path.join(STORAGE_DIR, output_filename)

    doc = Document()
    set_a4_margins(doc)

    # Atur gaya Normal ke Times New Roman 12pt spasi 1.5
    normal_style = doc.styles["Normal"]
    normal_font = normal_style.font
    normal_font.name = "Times New Roman"
    normal_font.size = Pt(12)
    normal_font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.space_after = Pt(6)

    # Judul Dokumen
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(title.upper())
    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(14)
    title_p.paragraph_format.space_after = Pt(24)

    # Identitas mahasiswa di bawah judul
    for idx, line in enumerate(identity_lines or []):
        id_p = doc.add_paragraph()
        id_p.paragraph_format.space_after = Pt(18 if idx == len(identity_lines) - 1 else 0)
        id_run = id_p.add_run(line)
        id_run.font.name = "Times New Roman"
        id_run.font.size = Pt(12)

    # Subbab dan Isi Naskah
    for section in sections:
        heading_text = section.get("heading", "")
        content_text = section.get("content", "")

        # Heading bab jika ada
        if heading_text and heading_text.strip():
            h_p = doc.add_paragraph()
            h_p.paragraph_format.space_before = Pt(14)
            h_p.paragraph_format.space_after = Pt(6)
            h_p.paragraph_format.keep_with_next = True
            h_run = h_p.add_run(heading_text)
            h_run.bold = True
            h_run.font.name = "Times New Roman"
            h_run.font.size = Pt(12)

        # Paragraf isi
        paragraphs = content_text.split("\n\n")
        for para in paragraphs:
            para_clean = para.strip()
            if not para_clean:
                continue
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Inches(0.5)
            p_run = p.add_run(para_clean)
            p_run.font.name = "Times New Roman"
            p_run.font.size = Pt(12)

    # Daftar Pustaka / References
    if language:
        is_en_doc = language.lower().startswith("en")
    else:
        sample_text = (title + " " + " ".join(s.get("content", "")[:200] for s in sections[:2])).lower()
        is_en_doc = any(w in sample_text.split() for w in ["the", "and", "is", "of", "to", "in", "that", "this", "urban", "living"])
    ref_title_text = "REFERENCES" if is_en_doc else "DAFTAR PUSTAKA"

    if not references:
        doc.save(file_path)
        return file_path

    ref_heading = doc.add_paragraph()
    ref_heading.paragraph_format.space_before = Pt(20)
    ref_heading.paragraph_format.space_after = Pt(10)
    ref_heading.paragraph_format.keep_with_next = True
    ref_h_run = ref_heading.add_run(ref_title_text)
    ref_h_run.bold = True
    ref_h_run.font.name = "Times New Roman"
    ref_h_run.font.size = Pt(12)

    for ref in references:
        ref_data = build_academic_reference_data(ref, is_en=is_en_doc)

        ref_p = doc.add_paragraph()
        ref_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        ref_p.paragraph_format.left_indent = Inches(0.5)
        ref_p.paragraph_format.first_line_indent = Inches(-0.5)
        ref_p.paragraph_format.space_after = Pt(6)

        # 1. Penulis dan Tahun
        lead_text = f"{ref_data['authors'].rstrip('.')}. ({ref_data['year']}). "
        run_lead = ref_p.add_run(lead_text)
        run_lead.font.name = "Times New Roman"
        run_lead.font.size = Pt(11)

        # 2. Judul Naskah
        run_title = ref_p.add_run(f"{ref_data['title']}.")
        run_title.font.name = "Times New Roman"
        run_title.font.size = Pt(11)

        # 3. Nama Jurnal / Publikasi (Cetak miring sesuai standar APA)
        if ref_data["venue"]:
            run_venue = ref_p.add_run(f" {ref_data['venue']}")
            run_venue.font.name = "Times New Roman"
            run_venue.font.size = Pt(11)
            run_venue.italic = True

        # 4. Detail Volume, Isu, Halaman
        if ref_data["pub_details"]:
            sep = ", " if ref_data["venue"] else " "
            run_details = ref_p.add_run(f"{sep}{ref_data['pub_details']}.")
            run_details.font.name = "Times New Roman"
            run_details.font.size = Pt(11)
        elif ref_data["venue"]:
            run_dot = ref_p.add_run(".")
            run_dot.font.name = "Times New Roman"
            run_dot.font.size = Pt(11)

        # 5. Tautan DOI
        if ref_data["doi"]:
            run_doi = ref_p.add_run(f" {ref_data['doi']}")
            run_doi.font.name = "Times New Roman"
            run_doi.font.size = Pt(11)

    doc.save(file_path)
    return file_path
