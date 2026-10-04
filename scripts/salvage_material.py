import json
import os
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def main():
    json_path = os.path.expandvars(r"%LOCALAPPDATA%\AsistenTugasCitra\storage\tasks\194ea1be.json")
    with open(json_path, "r", encoding="utf-8") as fh:
        data = json.load(fh)

    ref = data["references"][0]
    title = ref["title"]
    authors = ", ".join(ref["authors"])
    venue = ref["venue"]
    year = ref["year"]
    chunks = ref["pages_content"]

    downloads_dir = r"C:\Users\Galang\Downloads"
    docx_out = os.path.join(downloads_dir, "FSSI4105 Pengantar Linguistik Umum - Modul 1-3.docx")
    txt_out = os.path.join(downloads_dir, "FSSI4105 Pengantar Linguistik Umum - Modul 1-3.txt")

    # 1. Tulis TXT lengkap
    with open(txt_out, "w", encoding="utf-8") as f:
        f.write(f"{title}\n")
        f.write(f"Penulis: {authors}\n")
        f.write(f"Penerbit: {venue} ({year})\n")
        f.write("=" * 60 + "\n\n")
        for chunk in chunks:
            label = chunk.get("page_number", "")
            f.write(f"--- {label} ---\n\n")
            f.write(chunk.get("text", "").strip() + "\n\n")

    # 2. Tulis DOCX A4 yang rapi dan terformat
    doc = docx.Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    h1 = doc.add_heading(title, level=0)
    h1.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_meta.add_run(f"Penulis: {authors} | Penerbit: {venue} ({year})")
    r_meta.font.size = Pt(10)
    r_meta.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_paragraph()

    for chunk in chunks:
        label = chunk.get("page_number", "")
        doc.add_heading(label, level=2)
        paragraphs = chunk.get("text", "").split("\n\n")
        for p_text in paragraphs:
            clean_p = p_text.strip()
            if clean_p:
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(6)
                p.add_run(clean_p)

    doc.save(docx_out)
    print(f"Berhasil membuat DOCX: {docx_out} ({os.path.getsize(docx_out):,} bytes)")
    print(f"Berhasil membuat TXT: {txt_out} ({os.path.getsize(txt_out):,} bytes)")

    # 3. Simpan juga ke storage/modules untuk Bank Modul Lokal di repo
    modules_dir = os.path.abspath("storage/modules")
    os.makedirs(modules_dir, exist_ok=True)
    repo_json = os.path.join(modules_dir, "FSSI4105_Pengantar_Linguistik_Umum.json")
    with open(repo_json, "w", encoding="utf-8") as f:
        json.dump(ref, f, ensure_ascii=False, indent=2)
    print(f"Berhasil membuat modul backup JSON: {repo_json}")

if __name__ == "__main__":
    main()
