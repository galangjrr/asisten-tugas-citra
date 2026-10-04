import os
import pytest
from docx import Document
from pypdf import PdfReader
from fastapi.testclient import TestClient
from api.main import app
from api import routes

client = TestClient(app)


def test_endpoint_references_indonesian(monkeypatch):
    """
    Menguji endpoint /api/generate sampai pembuatan berkas docx dan pdf
    dengan format rujukan DAFTAR PUSTAKA standar APA 7th dalam bahasa Indonesia.
    """
    # 1. Daftarkan bahan bacaan ke cache
    mod_id_1 = "ref_jurnal_id"
    mod_id_2 = "ref_bmp_id"
    
    routes.CACHED_PAPERS[mod_id_1] = {
        "id": mod_id_1,
        "title": "Pengaruh Kecerdasan Buatan Terhadap Kinerja Belajar Mahasiswa",
        "authors": ["Budi Santoso", "Siti Aminah"],
        "year": 2022,
        "venue": "Jurnal Teknologi Pendidikan Indonesia",
        "volume": "14",
        "issue": "2",
        "pages": "115-128",
        "doi": "10.5555/jtpi.2022.14.2",
        "pdf_url": "",
        "all_pdf_urls": [],
        "abstract": "Kajian empiris mengenai pemanfaatan AI dalam proses belajar mengajar mahasiswa.",
        "scholar_url": "",
        "has_full_pdf": True,
        "is_manual_module": True,
        "is_ut_bmp": False,
        "source_status": "Bahan bacaan dosen",
        "pages_content": [{"page_number": 1, "text": "Isi rujukan jurnal AI."}],
    }

    routes.CACHED_PAPERS[mod_id_2] = {
        "id": mod_id_2,
        "title": "Manajemen (EKMA4116)",
        "authors": ["Universitas Terbuka"],
        "year": 2020,
        "venue": "Penerbit Universitas Terbuka",
        "volume": "",
        "issue": "",
        "pages": "",
        "page_info": "Modul 4, hlm. 10-22",
        "doi": "",
        "pdf_url": "",
        "all_pdf_urls": [],
        "abstract": "Modul dasar manajemen organisasi publik dan bisnis.",
        "scholar_url": "https://pustaka.ut.ac.id/",
        "has_full_pdf": True,
        "is_manual_module": True,
        "is_ut_bmp": True,
        "source_status": "Buku Materi Pokok UT",
        "pages_content": [{"page_number": "Modul 4, hlm. 10", "text": "Isi modul manajemen."}],
    }

    # 2. Mock pemanggilan agen generator Gemini
    async def fake_generator(**kwargs):
        return {
            "title": "Analisis Penerapan AI dalam Manajemen Pendidikan",
            "language": "id",
            "sections": [
                {
                    "heading": "1. Analisis Efektivitas AI",
                    "content": "Pemanfaatan AI terbukti meningkatkan efisiensi belajar mahasiswa secara signifikan."
                },
                {
                    "heading": "2. Integrasi Manajemen",
                    "content": "Tata kelola pembelajaran memerlukan perencanaan strategis organisasi pendidikan."
                }
            ]
        }

    monkeypatch.setattr(routes, "generate_academic_draft", fake_generator)

    # 3. Eksekusi endpoint /api/generate
    payload = {
        "topic": "Jelaskan peran AI dan manajemen dalam pendidikan tinggi.",
        "paper_ids": [mod_id_1, mod_id_2],
        "format_type": "jawaban_nomor",
        "target_words": 500,
        "paragraph_depth": "mendalam",
        "tone": "kritis",
        "student_name": "Galang Pratama",
        "student_id": "042918231",
        "course_name": "Manajemen Pendidikan",
    }

    response = client.post("/api/generate", json=payload)
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    task_id = data["task_id"]

    # 4. Unduh berkas melalui endpoint /api/download
    res_docx = client.get(f"/api/download/docx/{task_id}")
    assert res_docx.status_code == 200
    res_pdf = client.get(f"/api/download/pdf/{task_id}")
    assert res_pdf.status_code == 200

    # 5. Baca teks naskah docx dan verifikasi rujukan
    task = routes.TASKS_DB[task_id]
    doc = Document(task["docx_path"])
    docx_paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    # Temukan bagian DAFTAR PUSTAKA
    assert "DAFTAR PUSTAKA" in docx_paras
    ref_idx = docx_paras.index("DAFTAR PUSTAKA")
    ref_items = docx_paras[ref_idx + 1:]

    print("\n" + "=" * 60)
    print("OUTPUT HASIL GENERATE ENDPOINT: DAFTAR PUSTAKA (INDONESIA)")
    print("=" * 60)
    print(f"Judul Tugas: {data['title']}")
    print(f"Heading Referensi: {docx_paras[ref_idx]}")
    for i, r in enumerate(ref_items, 1):
        print(f"[{i}] {r}")
    print("=" * 60)

    # Verifikasi isi referensi
    assert any("Santoso, B., & Aminah, S. (2022)." in r for r in ref_items)
    assert any("Jurnal Teknologi Pendidikan Indonesia, 14(2), 115-128." in r for r in ref_items)
    assert any("https://doi.org/10.5555/jtpi.2022.14.2" in r for r in ref_items)
    assert any("Universitas Terbuka. (2020)." in r for r in ref_items)
    assert any("Manajemen (EKMA4116)." in r for r in ref_items)


def test_endpoint_references_english(monkeypatch):
    """
    Menguji endpoint /api/generate sampai pembuatan berkas docx dan pdf
    dengan format rujukan REFERENCES standar APA 7th dalam bahasa Inggris.
    """
    mod_id_en = "ref_en_journal"
    routes.CACHED_PAPERS[mod_id_en] = {
        "id": mod_id_en,
        "title": "Economic Growth in Metropolitan Regions",
        "authors": ["John Maynard Keynes", "Jane Smith"],
        "year": 2023,
        "venue": "Journal of Economic Perspectives",
        "volume": "37",
        "issue": "1",
        "pages": "45-60",
        "doi": "10.1257/jep.2023.37.1",
        "pdf_url": "",
        "all_pdf_urls": [],
        "abstract": "An empirical analysis of urban agglomeration and economic output.",
        "scholar_url": "",
        "has_full_pdf": True,
        "is_manual_module": True,
        "is_ut_bmp": False,
        "source_status": "Bahan bacaan dosen",
        "pages_content": [{"page_number": 1, "text": "Metropolitan economics text."}],
    }

    async def fake_generator(**kwargs):
        return {
            "title": "Metropolitan Economic Growth Analysis",
            "language": "en",
            "sections": [
                {
                    "heading": "1. Agglomeration Economies",
                    "content": "Urban concentration yields positive externalities through knowledge spillovers."
                }
            ]
        }

    monkeypatch.setattr(routes, "generate_academic_draft", fake_generator)

    payload = {
        "topic": "Explain the role of metropolitan economies in national economic growth.",
        "paper_ids": [mod_id_en],
        "format_type": "diskusi_mengalir",
        "target_words": 400,
        "paragraph_depth": "mendalam",
        "tone": "kritis",
        "student_name": "John Doe",
        "student_id": "098765432",
        "course_name": "Urban Economics",
    }

    response = client.post("/api/generate", json=payload)
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    task_id = data["task_id"]

    task = routes.TASKS_DB[task_id]
    doc = Document(task["docx_path"])
    docx_paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    assert "REFERENCES" in docx_paras
    assert "DAFTAR PUSTAKA" not in docx_paras
    ref_idx = docx_paras.index("REFERENCES")
    ref_items = docx_paras[ref_idx + 1:]

    print("\n" + "=" * 60)
    print("OUTPUT HASIL GENERATE ENDPOINT: REFERENCES (ENGLISH)")
    print("=" * 60)
    print(f"Judul Tugas: {data['title']}")
    print(f"Heading Referensi: {docx_paras[ref_idx]}")
    for i, r in enumerate(ref_items, 1):
        print(f"[{i}] {r}")
    print("=" * 60)

    assert any("Keynes, J. M., & Smith, J. (2023)." in r for r in ref_items)
    assert any("Journal of Economic Perspectives, 37(1), 45-60." in r for r in ref_items)
    assert any("https://doi.org/10.1257/jep.2023.37.1" in r for r in ref_items)
