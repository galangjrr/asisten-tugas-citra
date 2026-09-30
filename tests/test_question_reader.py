import io
import pytest
from fastapi.testclient import TestClient
import docx
import pypdf
from api.main import app
from tools.question_reader import parse_question_document, extract_course_code_from_text

client = TestClient(app)


def test_extract_course_code():
    text1 = "Tugas 1 Mata Kuliah EKMA4116 Manajemen Semester 2024.1"
    assert extract_course_code_from_text(text1) == "EKMA4116"

    text2 = "Petunjuk Diskusi FSSI 4206 English Syntax"
    assert extract_course_code_from_text(text2) == "FSSI4206"

    text3 = "Tanpa kode disini"
    assert extract_course_code_from_text(text3) is None


@pytest.mark.anyio
async def test_parse_docx_file():
    # Buat dokumen docx in-memory
    doc = docx.Document()
    doc.add_paragraph("Selamat pagi rekan mahasiswa sekalian.")
    doc.add_paragraph("Pertanyaan: Jelaskan perbedaan fonetik dan fonologi menurut modul FSSI4206!")
    
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    docx_bytes = stream.read()

    res = await parse_question_document(docx_bytes, "soal_tugas.docx")
    assert res["success"] is True
    assert res["file_type"] == "docx"
    assert "Jelaskan perbedaan fonetik" in res["text"]
    assert res["detected_course_code"] == "FSSI4206"


def test_upload_question_file_endpoint():
    doc = docx.Document()
    doc.add_paragraph("Soal Ujian Akhir Semester: Analisis strategi pemasaran PT ABC.")
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)

    response = client.post(
        "/api/upload-question-file",
        files={"file": ("soal.docx", stream.read(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert "Analisis strategi pemasaran" in data["text"]


def test_split_questions_and_guidelines():
    from tools.question_reader import split_questions_and_guidelines
    sample_text = """Petunjuk Pengerjaan Tugas:
1. Gunakan modul 3 BMP EKMA4116.
2. Rubrik penilaian: analisis 50%, kesimpulan 20%.

Soal 1:
Jelaskan konsep segmentasi pasar menurut Kotler!
"""
    result = split_questions_and_guidelines(sample_text)
    assert "Petunjuk Pengerjaan Tugas" in result["guidelines"]
    assert "Rubrik penilaian" in result["guidelines"]
    assert "Soal 1:" in result["questions"]
    assert "Jelaskan konsep segmentasi pasar" in result["questions"]


@pytest.mark.anyio
async def test_bilingual_assignment_parsing():
    indo_text = """TUGAS 1 TUTON UNIVERSITAS TERBUKA
Mata Kuliah: Manajemen Keuangan (EKMA4213)
Selamat pagi rekan mahasiswa.
Petunjuk:
1. Kerjakan minimal 800 kata.
2. Rujuk Modul 3 BMP.
Soal:
1. Jelaskan struktur modal optimal menurut trade-off theory!"""

    res = await parse_question_document(indo_text.encode("utf-8"), "tugas.txt")
    assert res["success"] is True
    assert "trade-off theory" in res["questions"]
    assert res["detected_course_code"] == "EKMA4213"




def test_split_ut_tutorial_sheet_with_metadata_header():
    from tools.question_reader import split_questions_and_guidelines
    sample_text = """FSSI4106 / English for Translation
Tutorial Assignment I - Session 3
Program Studi : Sastra Inggris
Butir Soal No. : 1 dan 2
Skor Maks : 100
Capaian Pembelajaran
Mahasiswa mampu menerjemahkan percakapan tertulis secara wajar.
Petunjuk:
• Terjemahkan setiap percakapan secara utuh ke dalam bahasa Indonesia.

FSSI4106 / English for Translation
Tutorial Assignment I - Session 3
SOAL TUGAS TUTORIAL I
1. Translate the following written conversation into Indonesian! (Score 50)
Guest Good evening. I have a reservation.
Total Score: 100"""
    result = split_questions_and_guidelines(sample_text)
    assert "Translate the following" in result["questions"]
    assert "Guest Good evening" in result["questions"]
    for noise in ("Program Studi", "Butir Soal No.", "Petunjuk", "Capaian", "English for Translation", "Total Score"):
        assert noise not in result["questions"]
    assert "Terjemahkan setiap percakapan" in result["guidelines"]
    assert "Capaian Pembelajaran" in result["guidelines"]
    assert "Total Score: 100" in result["guidelines"]
    assert "English for Translation" not in result["guidelines"]


def test_guess_answer_spec_for_translation_sheet():
    from tools.question_reader import guess_answer_spec
    questions = """1. Translate the following written conversation into Indonesian! (Score 50)
Guest Good evening.
2. Translate the following written conversation into Indonesian! (Score 50)
Host Welcome."""
    spec = guess_answer_spec(questions, "Petunjuk: Terjemahkan setiap percakapan secara utuh.")
    assert spec == {
        "question_count": 2,
        "answer_type": "terjemahan",
        "needs_citations": False,
        "answer_language": "id",
        "word_limit": None,
        "item_word_limits": None,
        "required_sections": None,
    }


def test_guess_answer_spec_for_essay_with_citations():
    from tools.question_reader import guess_answer_spec
    spec = guess_answer_spec("Write an essay about urban living.", "Use at least two references. Maximum 250-300 words.")
    assert spec["answer_type"] == "esai"
    assert spec["needs_citations"] is True
    assert spec["word_limit"] == 300
    assert spec["question_count"] is None


def test_normalize_answer_spec_rejects_invalid_ai_values():
    from tools.question_reader import normalize_answer_spec
    raw = {"question_count": "999", "answer_type": "puisi", "needs_citations": "ya", "answer_language": "fr", "word_limit": 500}
    spec = normalize_answer_spec(raw, "1. Jelaskan A\n2. Jelaskan B")
    assert spec["question_count"] == 2
    assert spec["answer_type"] == "jawaban_bernomor"
    assert spec["needs_citations"] is None
    assert spec["answer_language"] is None
    assert spec["word_limit"] == 500


def _numbered_paragraph(doc, text, num_id):
    """Paragraf dengan penomoran otomatis Word langsung di paragraf, seperti hasil tombol Numbering."""
    import copy
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    p = doc.add_paragraph(text)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), num_id)
    num_pr.append(ilvl)
    num_pr.append(num)
    p._p.get_or_add_pPr().append(copy.deepcopy(num_pr))
    return p


def test_read_docx_restores_list_numbers_in_document_order():
    from docx.oxml.ns import qn
    from tools.question_reader import read_docx_file, count_numbered_questions
    doc = docx.Document()
    num_id = doc.styles["List Number"].element.pPr.find(qn("w:numPr")).find(qn("w:numId")).get(qn("w:val"))

    doc.add_paragraph("LEMBAR SOAL TUGAS TUTORIAL I")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Indikator\t:\tMahasiswa mampu menemukan informasi dalam teks."
    doc.add_paragraph("Read the following questions. Then, scan the passage to find the answers.")
    _numbered_paragraph(doc, "Who built the tower?", num_id)
    _numbered_paragraph(doc, "How many people live there?", num_id)
    doc.add_paragraph("The tower was begun in 1066 by William the Conqueror.")

    stream = io.BytesIO()
    doc.save(stream)
    text = read_docx_file(stream.getvalue())

    # Tabel dibaca di posisinya, bukan ditumpuk di akhir dokumen
    assert text.index("Indikator") < text.index("Read the following")
    assert "1. Who built the tower?" in text
    assert "2. How many people live there?" in text
    assert count_numbered_questions(text) == 2


def test_count_numbered_questions_multi_part_sheet_is_not_locked():
    from tools.question_reader import count_numbered_questions
    sheet = "Answer the questions.\n1. Who?\n2. Where?\nTrue or false.\n1. The tower is old.\n2. The tower is new."
    assert count_numbered_questions(sheet) == 0


def test_read_docx_continues_numbering_from_start_override():
    import copy
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from tools.question_reader import read_docx_file, count_numbered_questions
    doc = docx.Document()
    base_id = doc.styles["List Number"].element.pPr.find(qn("w:numPr")).find(qn("w:numId")).get(qn("w:val"))

    # List kedua memakai startOverride 3, seperti bagian benar salah yang melanjutkan nomor soal 1-2
    numbering = doc.part.numbering_part.element
    base_num = next(n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == base_id)
    second = copy.deepcopy(base_num)
    second.set(qn("w:numId"), "99")
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "3")
    override.append(start)
    second.append(override)
    numbering.append(second)

    doc.add_paragraph("Answer the questions based on the passage.")
    _numbered_paragraph(doc, "Who built it?", base_id)
    _numbered_paragraph(doc, "When was it built?", base_id)
    doc.add_paragraph("Decide whether each statement is True (T) or False (F).")
    _numbered_paragraph(doc, "It is old.", "99")
    _numbered_paragraph(doc, "It is new.", "99")

    stream = io.BytesIO()
    doc.save(stream)
    text = read_docx_file(stream.getvalue())
    assert "3. It is old." in text
    assert "4. It is new." in text
    assert count_numbered_questions(text) == 4


def test_normalize_answer_spec_keeps_ai_answer_type():
    from tools.question_reader import normalize_answer_spec
    # Keputusan jenis jawaban milik AI, pola hanya mengisi yang kosong
    questions = "Read the passage.\n1. Who built the tower?\n2. Decide whether it is True (T) or False (F)."
    assert normalize_answer_spec({"answer_type": "jawaban_bernomor"}, questions)["answer_type"] == "jawaban_bernomor"
    assert normalize_answer_spec({"answer_type": None}, questions)["answer_type"] == "jawaban_singkat"


def test_guess_answer_spec_does_not_treat_generic_text_words_as_short_answer():
    from tools.question_reader import guess_answer_spec
    analysis = guess_answer_spec("1. Jelaskan konsep teks berikut dalam konteks manajemen.\n2. Bandingkan dengan teori Kotler.")
    assert analysis["answer_type"] == "jawaban_bernomor"

    # Permintaan referensi yang tertulis jelas tidak boleh dimatikan oleh kata 'bacaan'
    essay = guess_answer_spec("Setelah membaca bacaan pada Modul 3, analisis dampak kebijakan fiskal. Gunakan minimal 3 referensi.")
    assert essay["needs_citations"] is True
    assert essay["answer_type"] != "jawaban_singkat"


def test_strip_time_limit_lines_keeps_questions_that_mention_time():
    from tools.question_reader import strip_time_limit_lines
    sheet = "\n".join([
        "LEMBAR SOAL TUGAS TUTORIAL I",
        "Waktu\t:\t30 menit",
        "Time allotted: 2 hours",
        "Durasi 90 menit",
        "1. Berapa waktu tempuh kereta jika jaraknya 120 km?",
        "2. How many hours does the journey take?",
    ])
    cleaned = strip_time_limit_lines(sheet)
    assert "30 menit" not in cleaned
    assert "2 hours" not in cleaned
    assert "90 menit" not in cleaned
    # Soal yang kebetulan membahas waktu tetap utuh
    assert "1. Berapa waktu tempuh kereta" in cleaned
    assert "2. How many hours does the journey take?" in cleaned


def _make_pdf(pages):
    """PDF uji: item str jadi halaman teks biasa, item list jadi halaman gambar JPEG seperti hasil scan."""
    from PIL import Image, ImageDraw
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    stream = io.BytesIO()
    pdf = canvas.Canvas(stream, pagesize=A4)
    width, height = A4
    for page in pages:
        if isinstance(page, str):
            pdf.drawString(72, height - 72, page)
        else:
            image = Image.new("RGB", (1240, 1754), "white")
            draw = ImageDraw.Draw(image)
            for i, line in enumerate(page):
                draw.text((100, 120 + i * 60), line, fill="black")
            jpeg = io.BytesIO()
            image.save(jpeg, format="JPEG")
            jpeg.seek(0)
            pdf.drawImage(ImageReader(jpeg), 0, 0, width=width, height=height)
        pdf.showPage()
    pdf.save()
    return stream.getvalue()


@pytest.mark.anyio
async def test_read_pdf_ocr_runs_only_when_a_page_is_scanned(monkeypatch):
    import tools.question_reader as qr
    calls = []

    async def fake_ocr(pdf_bytes, mark_figures=False):
        calls.append(len(pdf_bytes))
        return "Halaman teks Soal nomor 1 lengkap.\n\n2. Soal dari halaman hasil scan yang dibaca OCR dengan lengkap."

    monkeypatch.setattr(qr, "extract_text_from_pdf", fake_ocr)

    text_only = _make_pdf(["1. Jelaskan konsep segmentasi pasar menurut Kotler secara lengkap."])
    assert "segmentasi pasar" in await qr.read_pdf_file(text_only)
    assert calls == []

    # Halaman kedua hasil scan, jadi seluruh PDF ditranskripsi walau halaman pertama berteks
    mixed = _make_pdf(["1. Jelaskan konsep segmentasi pasar menurut Kotler.", ["2. Soal dari halaman scan"]])
    result = await qr.read_pdf_file(mixed)
    assert len(calls) == 1
    assert "halaman hasil scan" in result


@pytest.mark.anyio
async def test_read_pdf_keeps_extracted_text_when_ocr_fails(monkeypatch):
    import tools.question_reader as qr

    async def failed_ocr(pdf_bytes, mark_figures=False):
        return ""

    monkeypatch.setattr(qr, "extract_text_from_pdf", failed_ocr)
    mixed = _make_pdf(["1. Jelaskan konsep segmentasi pasar menurut Kotler.", ["2. Soal scan"]])
    assert "segmentasi pasar" in await qr.read_pdf_file(mixed)


def _noise_png(size: int = 128) -> bytes:
    """PNG acak tanpa kompresi berarti supaya ukurannya melewati ambang gambar soal."""
    import os, struct, zlib
    raw = b"".join(b"\x00" + os.urandom(size * 3) for _ in range(size))

    def chunk(tag: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data))

    ihdr = struct.pack(">IIBBBBB", size, size, 8, 2, 0, 0, 0)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


@pytest.mark.anyio
async def test_read_docx_describes_question_images_in_place(monkeypatch):
    import tools.question_reader as qr

    async def fake_describe(blob, mime, question_context=""):
        assert mime == "image/png"
        # Pendeskripsi menerima perintah soal di sekitar gambar, di atas maupun di bawahnya
        assert "1. Describe the coffee shop" in question_context
        assert "2. Write a letter" in question_context
        assert "[GAMBAR INI]" in question_context
        return "Interior kafe dengan kursi merah dan meja bar kayu."

    monkeypatch.setattr(qr, "describe_image", fake_describe)
    doc = docx.Document()
    doc.add_paragraph("1. Describe the coffee shop in the picture below.")
    doc.add_picture(io.BytesIO(_noise_png()))
    doc.add_paragraph("2. Write a letter to the editor.")
    stream = io.BytesIO()
    doc.save(stream)

    text = await qr.read_docx_with_images(stream.getvalue())
    assert "[Gambar: Interior kafe dengan kursi merah" in text
    assert text.index("1. Describe") < text.index("[Gambar:") < text.index("2. Write")
    # Pembacaan tanpa list gambar tetap teks murni seperti sebelumnya
    assert "GAMBAR" not in qr.read_docx_file(stream.getvalue())


def test_read_docx_keeps_word_equations_as_linear_text():
    from docx.oxml import parse_xml
    from tools.question_reader import read_docx_file

    m = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'

    def r(t):
        return f"<m:r><m:t>{t}</m:t></m:r>"

    doc = docx.Document()
    p = doc.add_paragraph("1. Hitung ")
    p._p.append(parse_xml(
        f"<m:oMath {m}><m:f><m:num>{r('x+1')}</m:num><m:den>{r('2')}</m:den></m:f>{r('+')}"
        f"<m:sSup><m:e>{r('x')}</m:e><m:sup>{r('2')}</m:sup></m:sSup></m:oMath>"
    ))
    p.add_run(" jika x = 3.")
    p = doc.add_paragraph("2. Tentukan ")
    p._p.append(parse_xml(
        f"<m:oMath {m}><m:nary><m:naryPr><m:chr m:val='∑'/></m:naryPr><m:sub>{r('i=1')}</m:sub>"
        f"<m:sup>{r('n')}</m:sup><m:e>{r('i')}</m:e></m:nary>{r('+')}"
        f"<m:rad><m:deg/><m:e>{r('16')}</m:e></m:rad></m:oMath>"
    ))
    stream = io.BytesIO()
    doc.save(stream)

    text = read_docx_file(stream.getvalue())
    assert "1. Hitung (x+1)/2+x^2 jika x = 3." in text
    assert "2. Tentukan ∑_(i=1)^n i+sqrt(16)" in text


@pytest.mark.anyio
async def test_describe_image_drops_logos(monkeypatch):
    import tools.ocr_vision as ov

    class Res:
        text = "ABAIKAN"

    async def fake_generate(*args, **kwargs):
        return Res()

    monkeypatch.setattr(ov, "generate_with_fallback", fake_generate)
    assert await ov.describe_image(b"x" * 500) == ""


@pytest.mark.anyio
async def test_pdf_figures_are_cropped_and_described_in_order(monkeypatch):
    import tools.question_reader as qr
    marker_text = (
        "2. Tentukan matriks adjacency graf berikut.\n[[GAMBAR]]\n"
        "3. Diketahui graf di bawah ini.\n[[GAMBAR]]\na. Tentukan derajat tiap simpul."
    )
    contexts = []

    async def fake_ocr(pdf_bytes, mark_figures=False):
        return marker_text if mark_figures else "DESKRIPSI HALAMAN PENUH"

    async def fake_locate(page_png):
        return [(100, 100, 300, 500), (500, 100, 700, 500)]

    async def fake_describe(blob, mime, question_context=""):
        contexts.append(question_context)
        assert blob[1:4] == b"PNG"
        return f"graf ke-{len(contexts)}"

    monkeypatch.setattr(qr, "extract_text_from_pdf", fake_ocr)
    monkeypatch.setattr(qr, "locate_figures", fake_locate)
    monkeypatch.setattr(qr, "describe_image", fake_describe)
    pdf = _make_pdf(["Halaman soal graf"])

    text = await qr._transcribe_with_figures(pdf)
    assert text.index("[Gambar: graf ke-") < text.index("3. Diketahui") < text.rindex("[Gambar: graf ke-")
    # Setiap potongan menerima teks soal di sekitarnya dengan posisinya ditandai
    assert all("[GAMBAR INI]" in c for c in contexts)

    # Jumlah potongan tidak cocok dengan penanda, jadi urutan tidak dipercaya dan pakai deskripsi halaman penuh
    async def one_box(page_png):
        return [(100, 100, 300, 500)]

    monkeypatch.setattr(qr, "locate_figures", one_box)
    assert await qr._transcribe_with_figures(pdf) == "DESKRIPSI HALAMAN PENUH"


def test_item_word_limit_is_totalled_per_question():
    from tools.question_reader import guess_answer_spec, normalize_answer_spec

    questions = "1. Jelaskan konsep pasar.\n2. Jelaskan segmentasi.\n3. Jelaskan positioning."
    spec = guess_answer_spec(questions, "Jawaban maksimal 150 kata per soal. Waktu: 90 menit")
    assert spec["item_word_limits"] == [150, 150, 150]
    assert spec["word_limit"] == 450

    spec = guess_answer_spec(questions, "Tiap soal dijawab paling banyak 200 kata.")
    assert spec["item_word_limits"] == [200, 200, 200] and spec["word_limit"] == 600

    # Batas total tetap dibaca sebagai total, bukan dikali jumlah soal
    spec = guess_answer_spec(questions, "Seluruh jawaban maksimal 500 kata.")
    assert spec["word_limit"] == 500 and spec["item_word_limits"] is None

    # Batas berbeda per nomor dari AI dijumlahkan sendiri
    spec = normalize_answer_spec({"question_count": 2, "item_word_limits": [200, 300]}, "1. A\n2. B")
    assert spec["word_limit"] == 500 and spec["item_word_limits"] == [200, 300]

    # Tanpa batas di soal, keduanya kosong sehingga pengaturan panjang di form yang dipakai
    spec = guess_answer_spec(questions, "")
    assert spec["word_limit"] is None and spec["item_word_limits"] is None


def test_cut_figure_box_expands_to_whole_vector_graph():
    import math
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    import tools.question_reader as qr

    stream = io.BytesIO()
    c = canvas.Canvas(stream, pagesize=A4)
    c.drawString(60, 760, "1. Diketahui graf di bawah ini. Tentukan derajat tiap simpul.")
    pos = {v: (300 + 110 * math.cos(math.radians(90 - i * 60)), 560 + 110 * math.sin(math.radians(90 - i * 60))) for i, v in enumerate("ABCDEF")}
    for edge in ["A-B", "A-C", "A-D", "B-C", "B-E", "C-E", "D-E", "D-F", "E-F", "B-F"]:
        a, b = edge.split("-")
        c.line(*pos[a], *pos[b])
    for v, (x, y) in pos.items():
        c.circle(x, y, 6, fill=1)
        c.drawString(x + 10, y + 6, v)
    # Deretan pilihan pendek di bawah graf tidak boleh ikut masuk potongan
    for i, option in enumerate(["a. 10", "b. 12", "c. 14", "d. 16"]):
        c.drawString(290, 400 - i * 14, option)
    c.save()

    [(image, (drawings, labels))] = qr._render_pages(stream.getvalue())
    # Kotak dari Gemini yang hanya menangkap separuh kanan graf, seperti yang terjadi di uji nyata
    ymin, xmin, ymax, xmax = qr._expand_to_objects((183, 484, 481, 701), drawings, labels)
    width, height = A4
    left_vertex = min(x for x, _ in pos.values()) / width * 1000
    right_label = (max(x for x, _ in pos.values()) + 17) / width * 1000
    assert xmin <= left_vertex - 5 and xmax >= right_label
    # Pilihan jawaban berada di y 400 pt ke bawah, jadi batas bawah potongan harus di atasnya
    assert ymax < (height - 400 + 10) / height * 1000


def test_minimum_word_count_is_not_an_upper_limit():
    from tools.question_reader import extract_word_limit, normalize_answer_spec
    assert extract_word_limit("Write at least 200 words.") is None
    assert extract_word_limit("Write at least 200 words. Letter (200-250 words)") == 250
    assert extract_word_limit("minimal 300 kata") is None

    # Soal 1 cuma punya batas minimal, soal 2 punya rentang
    spec = normalize_answer_spec({"question_count": 2, "item_word_limits": [None, 250]}, "Describe the shop.\nWrite a letter.")
    assert spec["item_word_limits"] == [None, 250]
    assert spec["word_limit"] is None
    spec = normalize_answer_spec({"question_count": 2, "word_limit": 250, "item_word_limits": [None, 250]}, "Describe.\nWrite (200-250 words).")
    assert spec["word_limit"] is None


def test_extract_required_sections():
    from tools.question_reader import extract_required_sections, guess_answer_spec, normalize_answer_spec
    sheet = " Format Penulisan\n\nJudul\nPendahuluan\nPembahasan\nRefleksi\nKesimpulan\n"
    assert extract_required_sections(sheet) == ["Pendahuluan", "Pembahasan", "Refleksi", "Kesimpulan"]
    inline = "Sistematika penulisan: Judul, Pendahuluan, Isi dan Penutup. Daftar Pustaka"
    assert extract_required_sections(inline) == ["Pendahuluan", "Isi", "Penutup"]
    numbered = "Struktur makalah:\n1. Pendahuluan: latar belakang\n2. Pembahasan - analisis kasus\n3. Penutup\n\nKirim sebelum tanggal 5."
    assert extract_required_sections(numbered) == ["Pendahuluan", "Pembahasan", "Penutup"]
    # Aturan format huruf dan kalimat biasa bukan daftar bagian
    assert extract_required_sections("Format: Times New Roman 12, spasi 1.5") is None
    assert extract_required_sections("Format file PDF\nKirim lewat tuton\nJangan terlambat") is None
    assert guess_answer_spec(sheet)["required_sections"] == ["Pendahuluan", "Pembahasan", "Refleksi", "Kesimpulan"]
    # AI yang mengarang struktur dari rubrik kalah oleh pola teks, dan daftar AI tetap dibersihkan
    assert normalize_answer_spec({"required_sections": ["1. Pendahuluan", "Judul", "Isi"]}, "Jelaskan konsep X.")["required_sections"] == ["Pendahuluan", "Isi"]
    assert normalize_answer_spec({"required_sections": ["A"]}, sheet)["required_sections"] == ["Pendahuluan", "Pembahasan", "Refleksi", "Kesimpulan"]
